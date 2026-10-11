"""Builds a one-column, ATS-friendly .docx resume directly in Python (python-docx), no Node.js dependency required."""

from pathlib import Path
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .validate import load_profile
from .scoring import ScoredPosting


def _add_heading(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(12)
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    pPr = p._p.get_or_add_pPr()
    from docx.oxml.ns import qn
    from docx.oxml import OxmlElement
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "4")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "1F1F1F")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def _add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.add_run(text).font.size = Pt(10)
    return p


def build_resume_docx(scored: ScoredPosting, output_path: Path, profile: dict | None = None):
    profile = profile or load_profile()
    c = profile["candidate"]

    project_by_id = {p["id"]: p for p in profile["projects"]}
    skill_by_id = {s["id"]: s for s in profile["skills"]}

    skills_by_category: dict[str, list[str]] = {}
    for sid in scored.skill_ids:
        s = skill_by_id.get(sid)
        if not s:
            continue
        skills_by_category.setdefault(s["category"], []).append(s["name"])

    doc = Document()
    for section in doc.sections:
        section.top_margin = Pt(36)
        section.bottom_margin = Pt(36)
        section.left_margin = Pt(54)
        section.right_margin = Pt(54)

    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    title_p = doc.add_paragraph()
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title_p.add_run(c["name"])
    run.bold = True
    run.font.size = Pt(16)

    role_p = doc.add_paragraph()
    role_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role_run = role_p.add_run(f"Applicant \u2013 {scored.posting.title} ({scored.posting.company})")
    role_run.font.size = Pt(10.5)
    role_run.font.color.rgb = None

    contact_p = doc.add_paragraph()
    contact_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_p.add_run(f"{c['location']} | {c['email']}").font.size = Pt(9)

    links_p = doc.add_paragraph()
    links_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    links_p.add_run(
        f"{c['links']['github']} | {c['links']['portfolio']} | {c['links']['linkedin']}"
    ).font.size = Pt(9)

    _add_heading(doc, "EDUCATION")
    edu = c["education"]
    p = doc.add_paragraph()
    p.add_run(f"{edu['degree']} | {edu['school']}").bold = True
    p2 = doc.add_paragraph()
    p2.add_run(
        f"Expected Graduation: {edu['expected_graduation']} | GPA: {edu['gpa_4']}/4.0 ({edu['gpa_100']}/100)"
    ).font.size = Pt(10)

    _add_heading(doc, "TECHNICAL SKILLS")
    for category, names in skills_by_category.items():
        p = doc.add_paragraph()
        p.add_run(f"{category.title()}: ").bold = True
        p.add_run(", ".join(names))

    _add_heading(doc, "PROJECTS")
    for pid in scored.project_ids:
        proj = project_by_id.get(pid)
        if not proj:
            continue
        p = doc.add_paragraph()
        p.add_run(proj["title"]).bold = True
        meta_run = p.add_run(f"  |  {proj['meta']}")
        meta_run.italic = True
        meta_run.font.size = Pt(9)
        for bullet in proj["bullets"]:
            _add_bullet(doc, bullet)

    _add_heading(doc, "PROFESSIONAL EXPERIENCE")
    for exp in profile["experience"]:
        p = doc.add_paragraph()
        p.add_run(exp["title"]).bold = True
        meta_run = p.add_run(f"  |  {exp['meta']}")
        meta_run.italic = True
        meta_run.font.size = Pt(9)
        for bullet in exp["bullets"]:
            _add_bullet(doc, bullet)

    _add_heading(doc, "CERTIFICATIONS")
    for cert in profile["certifications"]:
        _add_bullet(doc, cert)

    _add_heading(doc, "LANGUAGES")
    doc.add_paragraph(" | ".join(f"{lang['language']}: {lang['level']}" for lang in c["languages"]))

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))
    return output_path


def convert_docx_to_pdf(docx_path: Path) -> Path | None:
    """Converts a .docx to .pdf in the same folder, using LibreOffice's headless mode."""
    import subprocess
    import shutil

    pdf_path = docx_path.with_suffix(".pdf")

    candidates = []
    on_path = shutil.which("soffice")
    if on_path:
        candidates.append(on_path)
    candidates += [
        r"C:\Program Files\LibreOffice\program\soffice.exe",
        r"C:\Program Files (x86)\LibreOffice\program\soffice.exe",
    ]

    for exe in candidates:
        if not exe or not Path(exe).exists():
            continue
        try:
            subprocess.run(
                [exe, "--headless", "--convert-to", "pdf", "--outdir", str(docx_path.parent), str(docx_path)],
                check=True,
                timeout=60,
                capture_output=True,
            )
        except (subprocess.CalledProcessError, subprocess.TimeoutExpired, OSError):
            continue
        if pdf_path.exists():
            return pdf_path

    return None
